import os
import shutil
import subprocess
import tempfile
from pathlib import Path

import pandas as pd
from bs4 import BeautifulSoup
from docx import Document as WordDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from langchain_core.documents import Document


class ExcelLoader:
    def __init__(self, file_path):
        self.file_path = file_path

    def load(self):
        documents = []
        sheets = pd.read_excel(self.file_path, sheet_name=None, dtype=str)
        for sheet_name, table in sheets.items():
            for row_number, row in table.iterrows():
                text = "\n".join(
                    f"{column}: {value}" for column, value in row.items() if pd.notna(value)
                )
                if text:
                    metadata = {"source": self.file_path, "sheet": sheet_name, "row": row_number}
                    documents.append(Document(page_content=text, metadata=metadata))
        return documents

class HTMLLoader:
    def __init__(self, file_path):
        self.file_path = file_path

    def load(self):
        with open(self.file_path, encoding="utf-8") as f:
            soup = BeautifulSoup(f.read(), "lxml")
        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()
        title = soup.title.string.strip() if soup.title and soup.title.string else ""
        text = soup.get_text(separator="\n", strip=True)
        return [Document(page_content=text, metadata={"source": self.file_path, "title": title})]


class DocxLoader:
    def __init__(self, file_path):
        self.file_path = file_path

    def load(self):
        word = WordDocument(self.file_path)
        lines = []
        for child in word.element.body.iterchildren():
            if child.tag.endswith("}p"):
                text = Paragraph(child, word).text.strip()
                if text:
                    lines.append(text)
            elif child.tag.endswith("}tbl"):
                for row in Table(child, word).rows:
                    lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        return [Document(page_content="\n".join(lines), metadata={"source": self.file_path})]


def find_libreoffice():
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    for candidate in (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ):
        if os.path.exists(candidate):
            return candidate
    raise RuntimeError("LibreOffice is needed to read .doc files but was not found")


class DocLoader:
    def __init__(self, file_path):
        self.file_path = file_path

    def load(self):
        soffice = find_libreoffice()
        stem = os.path.splitext(os.path.basename(self.file_path))[0]
        with tempfile.TemporaryDirectory() as out_dir, tempfile.TemporaryDirectory() as profile_dir:
            subprocess.run(
                [
                    soffice,
                    f"-env:UserInstallation={Path(profile_dir).as_uri()}",
                    "--headless",
                    "--convert-to",
                    "docx",
                    "--outdir",
                    out_dir,
                    self.file_path,
                ],
                check=True,
                capture_output=True,
                timeout=180,
            )
            documents = DocxLoader(os.path.join(out_dir, stem + ".docx")).load()
        for document in documents:
            document.metadata["source"] = self.file_path
        return documents
